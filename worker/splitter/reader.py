import pandas as pd
import numpy as np
import os

def read_file_preview(file_path: str, nrows: int = 20):
    """Reads file header, row count, columns, and sample rows with NaN/Inf sanitization."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.csv':
        df_iter = pd.read_csv(file_path, chunksize=nrows)
        first_chunk = next(df_iter)
        total_rows = sum(1 for _ in open(file_path, 'r', encoding='utf-8', errors='ignore')) - 1
        columns = list(first_chunk.columns)
        # Replace NaN / inf with None for JSON compliance
        preview_df = first_chunk.head(nrows).replace({np.nan: None, np.inf: None, -np.inf: None})
        preview_rows = preview_df.values.tolist()
        return {
            "columns": columns,
            "preview_rows": preview_rows,
            "total_rows": max(total_rows, 0),
            "detected_types": {col: str(first_chunk[col].dtype) for col in columns}
        }
    elif ext in ['.xlsx', '.xls']:
        df = pd.read_excel(file_path, nrows=nrows)
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        total_rows = sheet.max_row - 1 if sheet.max_row else 0
        columns = list(df.columns)
        
        preview_df = df.head(nrows).replace({np.nan: None, np.inf: None, -np.inf: None})
        preview_rows = preview_df.values.tolist()
        return {
            "columns": columns,
            "preview_rows": preview_rows,
            "total_rows": total_rows,
            "detected_types": {col: str(df[col].dtype) for col in columns}
        }
    else:
        raise ValueError("Unsupported file format. Use CSV or XLSX.")